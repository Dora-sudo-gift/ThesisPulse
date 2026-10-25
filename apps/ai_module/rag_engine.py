import os
import re
import fitz  
import docx
from apps.academic.models import Report


STOPWORDS = {
    "le", "la", "les", "un", "une", "des", "du", "de", "d", "et", "ou", "mais", "donc",
    "or", "ni", "car", "dans", "sur", "sous", "avec", "sans", "pour", "par", "en", "au",
    "aux", "ce", "cet", "cette", "ces", "mon", "ton", "son", "notre", "votre", "leur",
    "je", "tu", "il", "elle", "on", "nous", "vous", "ils", "elles", "qui", "que", "quoi",
    "dont", "où", "quand", "comment", "pourquoi", "quel", "quelle", "quels", "quelles",
    "est", "sont", "a", "ont", "fait", "faire", "dit", "dire", "peut", "plus", "moins",
    "très", "bien", "tout", "tous", "toute", "toutes", "comme", "être", "avoir", "été",
    "moi", "toi", "lui", "meme", "même", "autres", "autre", "entre", "vers", "chez",
    "the", "and", "for", "with", "this", "that", "from", "are", "was", "were", "have",
    "has", "had", "what", "which", "who", "whom", "whose", "where", "when", "why", "how",
    "dis", "moi", "parle", "texte", "document", "rapport", "projet", "etudiant", "étudiant",
    "partie", "concernant", "niveau", "sujet", "utilise", "utilisé", "utilisés", "utilisées"
}


TECH_KEYWORDS = [
    "Python", "Django", "Flask", "FastAPI", "Java", "Spring Boot", "JavaScript", "TypeScript",
    "React", "Angular", "Vue.js", "Next.js", "Node.js", "Express", "HTML", "CSS", "Tailwind",
    "Bootstrap", "PHP", "Laravel", "Symfony", "C#", ".NET", "C++", "Flutter", "React Native",
    "Kotlin", "Swift", "Android", "iOS", "PostgreSQL", "MySQL", "SQLite", "MongoDB", "Redis",
    "Oracle", "SQL Server", "Firebase", "Supabase", "Docker", "Kubernetes", "Docker Compose",
    "Jenkins", "GitLab CI", "GitHub Actions", "CI/CD", "Linux", "Ubuntu", "Kali Linux", "Nginx",
    "Apache", "AWS", "Azure", "Google Cloud", "GCP", "Microservices", "REST API", "GraphQL",
    "gRPC", "WebSockets", "JWT", "OAuth2", "OpenAI", "LLM", "RAG", "LangChain", "HuggingFace",
    "TensorFlow", "PyTorch", "Scikit-learn", "Pandas", "NumPy", "Matplotlib", "OR-Tools",
    "Wireshark", "Metasploit", "Nmap", "OWASP", "RFID", "IoT", "Arduino", "Raspberry Pi",
    "ESP32", "MQTT", "Blockchain", "Solidity", "Smart Contracts", "UML", "Merise", "Scrum", "Agile"
]


def extract_text_from_file(file_path):
    
    text = ""
    if not file_path or not os.path.exists(file_path):
        return text

    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext == '.pdf':
            try:
                with fitz.open(file_path) as doc:
                    for page_num, page in enumerate(doc, start=1):
                        page_text = page.get_text("text")
                        if page_text and page_text.strip():
                            text += f"\n[Page {page_num}]\n" + page_text + "\n"
            except Exception:
                try:
                    import pypdf
                    reader = pypdf.PdfReader(file_path)
                    for page_num, page in enumerate(reader.pages, start=1):
                        page_text = page.extract_text() or ""
                        if page_text.strip():
                            text += f"\n[Page {page_num}]\n" + page_text + "\n"
                except Exception:
                    import PyPDF2
                    with open(file_path, 'rb') as f:
                        reader = PyPDF2.PdfReader(f)
                        for page_num, page in enumerate(reader.pages, start=1):
                            page_text = page.extract_text() or ""
                            if page_text.strip():
                                text += f"\n[Page {page_num}]\n" + page_text + "\n"

        elif ext in ['.docx', '.doc']:
            doc = docx.Document(file_path)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        paragraphs.append(row_text)
            text = "\n\n".join(paragraphs)
    except Exception as e:
        print(f"Erreur lors de l'extraction du fichier {file_path}: {e}")

    return text.strip()


def _split_into_chunks(text):
    
    cleaned = re.sub(r'\r\n', '\n', text)
    raw_blocks = re.split(r'\n\s*\n', cleaned)
    
    chunks = []
    current_chunk = ""
    
    for block in raw_blocks:
        block = re.sub(r'\s+', ' ', block).strip()
        if not block:
            continue
        if len(current_chunk) + len(block) < 450:
            current_chunk = f"{current_chunk} {block}".strip()
        else:
            if len(current_chunk) > 30:
                chunks.append(current_chunk)
            current_chunk = block
            
    if len(current_chunk) > 20:
        chunks.append(current_chunk)


    if len(chunks) <= 1 and len(text) > 500:
        sentences = re.split(r'(?<=[.!?])\s+', re.sub(r'\s+', ' ', text))
        chunks = []
        temp = ""
        for s in sentences:
            if len(temp) + len(s) < 450:
                temp = f"{temp} {s}".strip()
            else:
                if temp:
                    chunks.append(temp)
                temp = s
        if temp:
            chunks.append(temp)

    return chunks


def _extract_detected_tags(text):

    found_tags = []
    text_lower = text.lower()
    for tech in TECH_KEYWORDS:
        
        pattern = r'\b' + re.escape(tech.lower()) + r'\b'
        if re.search(pattern, text_lower):
            found_tags.append(tech)
    return found_tags


def _extract_headings(text):
    
    headings = []
    for line in text.splitlines():
        line_clean = line.strip()
        if 4 < len(line_clean) < 95:
            if (
                re.match(r'^(chapitre|section|partie|module|étape|I+\.|[1-9]\.|[1-9]\.[1-9])\b', line_clean, re.IGNORECASE)
                or (line_clean.isupper() and len(line_clean.split()) >= 2)
                or line_clean.endswith(':')
            ):
                if line_clean not in headings:
                    headings.append(line_clean)
    return headings[:10]


def query_project_chatbot(project_id, user_question, report_id=None):
    
    reports = Report.objects.filter(project_id=project_id).order_by('-uploaded_at')
    if report_id:
        reports = reports.filter(id=report_id)

    if not reports.exists():
        return "❌ Aucun rapport PDF ou Word n'a encore été déposé pour ce projet."

    docs_data = []
    full_text = ""


    for rep in reports:
        if rep.file:
            extracted = extract_text_from_file(rep.file.path)
            if extracted:
                doc_name = os.path.basename(rep.file.name)
                docs_data.append({
                    'id': rep.id,
                    'name': doc_name,
                    'text': extracted
                })
                full_text += f"\n\n=== DOCUMENT #ID-{rep.id} ({doc_name}) ===\n{extracted}"

    if not full_text.strip():
        return "⚠️ Le fichier déposé semble vide ou contient uniquement des images scannées sans texte sélectionnable."

    chunks = _split_into_chunks(full_text)
    detected_tags = _extract_detected_tags(full_text)
    headings = _extract_headings(full_text)
    word_count = len(full_text.split())

    q_lower = user_question.lower().strip()
    
    
    raw_words = re.findall(r'[a-zA-ZÀ-ÿ0-9_-]{3,}', q_lower)
    query_keywords = [w for w in raw_words if w not in STOPWORDS]


    if any(k in q_lower for k in ["tag", "tags", "techno", "technologie", "technologies", "stack", "outil", "outils", "langage", "framework", "bibliothèque"]):
        tags_str = ", ".join(detected_tags) if detected_tags else "Aucune technologie standard détectée automatiquement."
        
        
        tech_Wait = []
        for c in chunks:
            if any(t.lower() in c.lower() for t in detected_tags[:6]) or any(w in c.lower() for w in ["technologie", "architecture", "outil", "environnement", "stack", "base de données", "framework"]):
                tech_Wait.append(c)
        
        excerpts = "\n\n".join([f"• « {p[:350]}... »" if len(p) > 350 else f"• « {p} »" for p in tech_Wait[:3]])
        if not excerpts and chunks:
            excerpts = f"• « {chunks[0][:300]}... »"

        return (
            f"🏷️ **Technologies, Tags & Outils identifiés dans le document** :\n"
            f"**{tags_str}**\n\n"
            f"📖 **Passages exacts du document qui les mentionnent** :\n"
            f"{excerpts}"
        )


    is_summary_request = any(k in q_lower for k in ["résume", "resume", "résumé", "resumer", "synthèse", "synthese", "globalement", "parle quoi", "dit quoi", "qu'est-ce qu'il a dit", "contenu du document"])
    
    
    if is_summary_request and len(query_keywords) <= 3:
        intro_snippet = chunks[0][:400] if chunks else ""
        middle_snippet = chunks[len(chunks) // 2][:350] if len(chunks) > 2 else ""
        end_snippet = chunks[-1][:350] if len(chunks) > 1 else ""

        headings_str = "\n".join([f"  - {h}" for h in headings[:6]]) if headings else "  - Document continu (sans titres de sections explicites)"
        tags_line = ", ".join(detected_tags[:10]) if detected_tags else "Non spécifiées explicitement"

        summary_parts = [
            f"📄 **Résumé fidèle du document ({docs_data[0]['name']} — {word_count} mots)** :",
            f"\n📌 **Structure / Points abordés dans le fichier** :\n{headings_str}",
            f"\n🛠️ **Mots-clés & Technologies cités** : {tags_line}",
            f"\n📝 **Ce que dit concrètement l'étudiant dans son document** :",
            f"• **Début / Introduction** : « {intro_snippet} »"
        ]
        if middle_snippet:
            summary_parts.append(f"• **Développement / Cœur du sujet** : « {middle_snippet} »")
        if end_snippet and end_snippet != intro_snippet:
            summary_parts.append(f"• **Fin / Conclusion du rapport** : « {end_snippet} »")

        return "\n".join(summary_parts)

  
    scored_chunks = []
    for chunk in chunks:
        chunk_lower = chunk.lower()
        score = 0
        for kw in query_keywords:
            
            occurrences = len(re.findall(r'\b' + re.escape(kw) + r'\b', chunk_lower))
            if occurrences > 0:
                score += occurrences * 10
            elif kw in chunk_lower:
                score += 4
        if score > 0:
            scored_chunks.append((score, chunk))

    scored_chunks.sort(key=lambda x: x[0], reverse=True)

    if scored_chunks:
        top_passages = [c for _, c in scored_chunks[:3]]
        formatted_passages = "\n\n".join([f"🔹 « {p} »" for p in top_passages])
        
        return (
            f"🎯 **Voici exactement ce que dit le document concernant votre question** (*\"{user_question}\"*) :\n\n"
            f"{formatted_passages}\n\n"
            f"📌 **Tags/Concepts liés dans le rapport** : {', '.join(detected_tags[:8]) if detected_tags else 'Aucun tag technique spécifique'}"
        )


    preview_passages = "\n\n".join([f"• « {c[:320]} »" for c in chunks[:3]])
    return (
        f"ℹ️ **Le document ne mentionne pas explicitement les termes `{', '.join(query_keywords) if query_keywords else user_question}`.**\n\n"
        f"Voici toutefois ce que contient réellement le document déposé (**{docs_data[0]['name']}**, {word_count} mots) :\n\n"
        f"{preview_passages}\n\n"
        f"🏷️ **Technologies détectées dans le fichier** : {', '.join(detected_tags) if detected_tags else 'Aucune'}"
    )